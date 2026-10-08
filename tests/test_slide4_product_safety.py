from __future__ import annotations

import io
import unittest
from pathlib import Path

from pptx import Presentation
from unittest.mock import patch

from app.audit_helpers.strategic_identity import _match_style_product_type, resolve_strategic_identity
from audit_powerpoint_new import (
    _build_slide4_column,
    build_slide4_pdp_benchmark_payload,
    generate_new_audit_powerpoint_from_template,
    _slide4_add_candidate,
    _slide4_evidence_theme,
    _slide4_product_language_guard,
    _slide4_record_facts,
)


def record(title: str, product_type: str, category: str = "Food & Beverage") -> dict:
    return {
        "product_title": title, "product_type": product_type, "category": category,
        "brand": "Test Brand", "description_body": "Product details and usage guidance.",
        "review_count": 120, "image_count": 8, "images": [],
    }


class Slide4ProductSafetyTests(unittest.TestCase):
    def test_fish_protein_and_recipes_do_not_establish_peanut_butter(self):
        fish = record("Wild Caught Salmon", "Canned & Jarred Seafood")
        fish.update({
            "description_body": "Protein and nutrition for pantry meals. Serve with peanut butter sauce or cocoa mole.",
            "key_features": ["Wild caught salmon", "Protein per serving"],
        })
        self.assertEqual(_slide4_evidence_theme(_slide4_record_facts(fish)), "product")
        for side in ("client", "competitor_1"):
            with self.subTest(side=side):
                column = _build_slide4_column(fish, "Fish", side=side)
                self.assertEqual(len(column["bullets"]), 6)
                text = " ".join(column["bullets"]).lower()
                for forbidden in ("peanut", "hazelnut", "almond butter", "breakfast", "snack and pantry"):
                    self.assertNotIn(forbidden, text)
                self.assertTrue(any("nutrition" in bullet.lower() for bullet in column["bullets"]))

    def test_recipe_ingredient_in_title_does_not_override_specific_product_type(self):
        fish = record("Salmon with Peanut Butter Sauce", "Canned & Jarred Seafood")
        self.assertEqual(_slide4_evidence_theme(_slide4_record_facts(fish)), "product")
        column = _build_slide4_column(fish, "Fish")
        self.assertNotIn("peanut butter", " ".join(column["bullets"]).lower())

    def test_unrelated_recipe_cannot_reclassify_unknown_product(self):
        fish = record("Smoked Salmon Fillet", "Unclassified")
        fish["description_body"] = "Serve with peanut butter, hazelnut spread, nutella, almond butter, and nut spread."
        identity = resolve_strategic_identity([fish])
        self.assertNotEqual(identity["product_type_display"], "Nut Butters & Spreads")
        self.assertNotEqual(identity["family_key"], "dips_spreads")
        column = _build_slide4_column(fish, "Fish")
        text = " ".join(column["bullets"]).lower()
        for forbidden in ("peanut", "hazelnut", "almond butter"):
            self.assertNotIn(forbidden, text)

    def test_exact_product_type_outweighs_description_matches(self):
        fish = record("Wild Caught Salmon", "Canned & Jarred Seafood")
        fish["description_body"] = "Creamy crunchy peanut butter almond butter cashew butter hazelnut spread nutella jif."
        identity = resolve_strategic_identity([fish])
        self.assertEqual(identity["product_type_display"], "Canned & Jarred Seafood")

    def test_same_category_guard_rejects_another_product_name(self):
        fish = record("Canned Salmon", "Canned & Jarred Seafood")
        facts = _slide4_record_facts(fish)
        guide_context = _slide4_product_language_guard(resolve_strategic_identity([fish]), facts)
        candidates, rejected = [], []
        _slide4_add_candidate(
            candidates, rejected, text="Peanut butter positioning reinforces pantry relevance",
            family="positioning_title", evidence_source="test", score=90,
            bullet_type="strength", dimension="product_positioning", signals=[], reason="Injected wrong product",
            product_context="canned and jarred seafood", category_context="food", guide_context=guide_context,
        )
        self.assertTrue(any("unsupported_product_identity" in item["reason"] for item in rejected))
        self.assertNotIn("peanut", " ".join(item["text"] for item in candidates).lower())

    def test_product_name_appears_once_across_different_categories(self):
        cases = (
            ("Canned Salmon", "Canned & Jarred Seafood", "Food & Beverage"),
            ("Daily Facial Cleanser", "Facial Cleansers", "Beauty"),
            ("Squishy Toys Set", "Squishy Toys", "Toys"),
            ("Wireless Headphones", "Headphones", "Electronics"),
            ("Unknown Product", "Unclassified", "Everything Else"),
        )
        for title, product_type, category in cases:
            with self.subTest(product_type=product_type):
                item = record(title, product_type, category)
                column = _build_slide4_column(item, "Client")
                self.assertEqual(len(column["bullets"]), 6)
                self.assertEqual(len(set(column["bullets"])), 6)
                product = column["bullet_debug"][0]["product_type_context"].lower()
                if product not in {"product", "the product"}:
                    self.assertLessEqual(" ".join(column["bullets"]).lower().count(product), 1)
                self.assertEqual([d["text"] for d in column["bullet_debug"]], column["bullets"])

    def test_sparse_evidence_uses_guarded_neutral_recommendations(self):
        item = {"product_title": "Unknown Widget", "category": "Electronics", "product_type": "Unclassified"}
        with patch("audit_powerpoint_new._translate_slide4_strategic_cues", return_value=([], [], {})):
            column = _build_slide4_column(item, "Client")
        self.assertEqual(len(column["bullets"]), 6)
        self.assertTrue(any(d["evidence_family_source"] == "controlled_fallback" for d in column["bullet_debug"]))
        text = " ".join(column["bullets"]).lower()
        for forbidden in ("peanut", "recipe", "nutrition", "breakfast"):
            self.assertNotIn(forbidden, text)

    def test_real_spreads_keep_supported_product_wording(self):
        for title, expected in (("Creamy Peanut Butter", "peanut"), ("Nutella Hazelnut Spread with Cocoa", "hazelnut")):
            with self.subTest(title=title):
                item = record(title, "Nut Butters & Spreads")
                item["description_body"] = "Breakfast and snack recipe suggestions with nutrition and protein details."
                column = _build_slide4_column(item, "Client")
                self.assertEqual(len(column["bullets"]), 6)
                self.assertIn(expected, " ".join(column["bullets"]).lower())

    @unittest.skipUnless((Path(__file__).resolve().parents[1] / "templates" / "Audit_Template_New.pptx").exists(), "Template unavailable")
    def test_fish_bullets_survive_powerpoint_export_without_wrong_product_or_repetition(self):
        client = record("Canned Salmon", "Canned & Jarred Seafood")
        client["brand"] = "Fish Client"
        competitor = record("Canned Salmon", "Canned & Jarred Seafood")
        competitor.update({"brand": "Wixter Seafood", "description_body": "Protein, nutrition, and pantry serving suggestions."})
        plan = {
            "audit_metadata": {"client_name": "Fish Client", "audit_date": "2026-10-08"},
            "product_slide_pairs": [{"pdp_slide": client}],
        }
        payload = build_slide4_pdp_benchmark_payload(plan, competitor_records=[competitor])
        bullets = [text for column in payload["columns"] if column["active"] for text in column["bullets"]]
        self.assertEqual(len(bullets), 12)
        self.assertEqual(len(set(bullets)), 12)
        deck = generate_new_audit_powerpoint_from_template(
            export_plan=plan,
            template_path=str(Path(__file__).resolve().parents[1] / "templates" / "Audit_Template_New.pptx"),
            competitor_records=[competitor],
        )
        presentation = Presentation(io.BytesIO(deck))
        slide = next(slide for slide in presentation.slides if any(
            "PDP Content Benchmarking" in (getattr(shape, "text", "") or "") for shape in slide.shapes
        ))
        text = " ".join(getattr(shape, "text", "") for shape in slide.shapes).lower()
        self.assertNotIn("peanut", text)
        self.assertNotIn("breakfast", text)
        for bullet in bullets:
            self.assertIn(bullet.lower(), text)
        for column in payload["columns"]:
            if column["active"]:
                self.assertLessEqual(" ".join(column["bullets"]).lower().count("canned and jarred seafood"), 1)

    def test_equal_product_matches_keep_supplied_broad_label(self):
        guide = {"families": {"family": {"product_types": {
            "first": {"product_type_display": "Type One", "aliases": ["widget"]},
            "second": {"product_type_display": "Type Two", "aliases": ["widget"]},
        }}}}
        matched = _match_style_product_type(guide, "Unclassified", "widget", identity_text="widget")
        self.assertEqual(matched[0], "")
        self.assertEqual(matched[1], "Unclassified")


if __name__ == "__main__":
    unittest.main()
